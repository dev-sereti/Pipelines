import openmeteo_requests
import pandas as pd
import requests_cache
from retry_requests import retry
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from pathlib import Path

# Output directory

OUTPUT_DIR = Path(r"C:\Users\Sereti\OneDrive - Victory Farms Ltd\Documents\Weather")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

EXCEL_FILE = OUTPUT_DIR / "weather_forecast.xlsx"

# Location mapping (coordinates must match the order in params exactly)

LOCATIONS = [
    {"name": "Roo Farm", "lat": -0.5603, "lon": 34.0623},
    {"name": "Kagano",   "lat": -2.3328, "lon": 29.0934},
    {"name": "Kigembe",  "lat": -2.7334, "lon": 29.7245},
]

LOCATION_COLORS = {
    "Roo Farm": "D6EAF8",
    "Kagano":   "D5F5E3",
    "Kigembe":  "FCF3CF",
}

COLUMN_HEADERS = [
    "Date & Time",
    "Location",
    "Latitude (°N)",
    "Longitude (°E)",
    "Temperature 180m (°C)",
    "Wind Gusts 10m (km/h)",
    "Wind Direction 180m (°)",
    "Wind Speed 180m (km/h)",
    "Rain (mm)",
    "Relative Humidity 2m (%)",
]
# Column rename map(Raw API names to Excel header names)
RENAME_MAP = {
    "date":                 "Date & Time",
    "location":             "Location",
    "latitude":             "Latitude (°N)",
    "longitude":            "Longitude (°E)",
    "temperature_180m":     "Temperature 180m (°C)",
    "wind_gusts_10m":       "Wind Gusts 10m (km/h)",
    "wind_direction_180m":  "Wind Direction 180m (°)",
    "wind_speed_180m":      "Wind Speed 180m (km/h)",
    "rain":                 "Rain (mm)",
    "relative_humidity_2m": "Relative Humidity 2m (%)",
}


# API Setup

cache_session = requests_cache.CachedSession('.cache', expire_after=3600)
retry_session = retry(cache_session, retries=5, backoff_factor=0.2)
openmeteo     = openmeteo_requests.Client(session=retry_session)

url = "https://api.open-meteo.com/v1/forecast"
params = {
    "latitude":  [-0.5603, -2.3328, -2.7334],
    "longitude": [34.0623, 29.0934, 29.7245],
    "hourly": [
        "temperature_180m",
        "wind_gusts_10m",
        "wind_direction_180m",
        "wind_speed_180m",
        "rain",
        "relative_humidity_2m",
    ],
    "timezone":  "auto",
    "past_days": 31,
}

responses = openmeteo.weather_api(url, params=params)

# Process 3 locations

all_frames = []

for i, response in enumerate(responses):
    print(f"\nCoordinates: {response.Latitude()}°N {response.Longitude()}°E")
    print(f"Elevation: {response.Elevation()} m asl")
    print(f"Timezone: {response.Timezone()}{response.TimezoneAbbreviation()}")
    print(f"Timezone difference to GMT+0: {response.UtcOffsetSeconds()}s")

    hourly = response.Hourly()
    hourly_temperature_180m     = hourly.Variables(0).ValuesAsNumpy()
    hourly_wind_gusts_10m       = hourly.Variables(1).ValuesAsNumpy()
    hourly_wind_direction_180m  = hourly.Variables(2).ValuesAsNumpy()
    hourly_wind_speed_180m      = hourly.Variables(3).ValuesAsNumpy()
    hourly_rain                 = hourly.Variables(4).ValuesAsNumpy()
    hourly_relative_humidity_2m = hourly.Variables(5).ValuesAsNumpy()

    hourly_data = {"date": pd.date_range(
        start     = pd.to_datetime(hourly.Time() + response.UtcOffsetSeconds(), unit="s", utc=True),
        end       = pd.to_datetime(hourly.TimeEnd() + response.UtcOffsetSeconds(), unit="s", utc=True),
        freq      = pd.Timedelta(seconds=hourly.Interval()),
        inclusive = "left"
    )}

    hourly_data["temperature_180m"]     = hourly_temperature_180m
    hourly_data["wind_gusts_10m"]       = hourly_wind_gusts_10m
    hourly_data["wind_direction_180m"]  = hourly_wind_direction_180m
    hourly_data["wind_speed_180m"]      = hourly_wind_speed_180m
    hourly_data["rain"]                 = hourly_rain
    hourly_data["relative_humidity_2m"] = hourly_relative_humidity_2m

    hourly_dataframe = pd.DataFrame(data=hourly_data)
    print("\nHourly data\n", hourly_dataframe)

    hourly_dataframe["location"]  = LOCATIONS[i]["name"]
    hourly_dataframe["latitude"]  = LOCATIONS[i]["lat"]
    hourly_dataframe["longitude"] = LOCATIONS[i]["lon"]
    hourly_dataframe["date"]      = hourly_dataframe["date"].dt.tz_localize(None)

    all_frames.append(hourly_dataframe)

# Combine all locations and rename columns
new_df = pd.concat(all_frames, ignore_index=True)
new_df = new_df.rename(columns=RENAME_MAP)
new_df["Date & Time"] = pd.to_datetime(new_df["Date & Time"], errors="coerce")


# Load existing data if file exists

def load_existing_data(path: Path) -> pd.DataFrame:
    if not path.exists():
        print("No existing file found — creating new file.")
        return pd.DataFrame(columns=COLUMN_HEADERS)

    try:
        existing = pd.read_excel(path, sheet_name="All Locations")
        existing["Date & Time"] = pd.to_datetime(existing["Date & Time"], errors="coerce")
        print(f"Loaded {len(existing)} existing rows from {path.name}")
        return existing
    except Exception as e:
        print(f"Could not read existing file ({e}) — creating new file.")
        return pd.DataFrame(columns=COLUMN_HEADERS)

# Merge and deduplicate

def merge_and_deduplicate(existing_df: pd.DataFrame, new_df: pd.DataFrame) -> pd.DataFrame:
    if existing_df.empty:
        print(f"First run — {len(new_df)} rows to save.")
        return new_df

    combined = pd.concat([existing_df, new_df], ignore_index=True)
    combined["Date & Time"] = pd.to_datetime(combined["Date & Time"], errors="coerce")
    combined.drop_duplicates(subset=["Date & Time", "Location"], keep="last", inplace=True)
    combined.sort_values(["Location", "Date & Time"], inplace=True)
    combined.reset_index(drop=True, inplace=True)

    new_rows = len(combined) - len(existing_df)
    print(f"Appended {new_rows} new rows — total: {len(combined)} rows")
    return combined

# Excel formatting helpers

def thin_border():
    s = Side(style="thin", color="BFBFBF")
    return Border(left=s, right=s, top=s, bottom=s)

def header_style(cell):
    cell.font      = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    cell.fill      = PatternFill("solid", start_color="1F4E79")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border    = thin_border()

def data_style(cell, even_row=False, loc_color=None):
    cell.font      = Font(name="Arial", size=10)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border    = thin_border()
    if loc_color:
        cell.fill = PatternFill("solid", start_color=loc_color)
    elif even_row:
        cell.fill = PatternFill("solid", start_color="EBF3FB")
    else:
        cell.fill = PatternFill("solid", start_color="FFFFFF")

# Save to Excel

def save_to_excel(df: pd.DataFrame, path: Path):
    wb = Workbook()

    #  All Locations sheet 
    ws = wb.active
    ws.title = "All Locations"
    ws.append(COLUMN_HEADERS)
    for cell in ws[1]:
        header_style(cell)

    for row_idx, row in df.iterrows():
        even      = row_idx % 2 == 0
        loc       = str(row["Location"])
        loc_color = LOCATION_COLORS.get(loc)
        dt_val    = row["Date & Time"]
        dt_str    = dt_val.strftime("%Y-%m-%d %H:%M") if pd.notna(dt_val) else ""

        ws.append([
            dt_str,
            loc,
            row["Latitude (°N)"],
            row["Longitude (°E)"],
            round(float(row["Temperature 180m (°C)"]),    2),
            round(float(row["Wind Gusts 10m (km/h)"]),    2),
            round(float(row["Wind Direction 180m (°)"]),  1),
            round(float(row["Wind Speed 180m (km/h)"]),   2),
            round(float(row["Rain (mm)"]),                3),
            round(float(row["Relative Humidity 2m (%)"]), 1),
        ])
        for col_idx, cell in enumerate(ws[ws.max_row], 1):
            if col_idx in (2, 3, 4):
                data_style(cell, loc_color=loc_color)
            else:
                data_style(cell, even_row=even)

    col_widths = [22, 14, 16, 17, 24, 22, 26, 24, 12, 26]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"

    #  Sheet per location 
    for loc in LOCATIONS:
        loc_df = df[df["Location"] == loc["name"]].copy()
        if loc_df.empty:
            continue

        ws_loc = wb.create_sheet(title=loc["name"])
        ws_loc.append(COLUMN_HEADERS)
        for cell in ws_loc[1]:
            header_style(cell)

        for row_idx, row in loc_df.iterrows():
            even   = row_idx % 2 == 0
            dt_val = row["Date & Time"]
            dt_str = dt_val.strftime("%Y-%m-%d %H:%M") if pd.notna(dt_val) else ""
            ws_loc.append([
                dt_str,
                str(row["Location"]),
                row["Latitude (°N)"],
                row["Longitude (°E)"],
                round(float(row["Temperature 180m (°C)"]),    2),
                round(float(row["Wind Gusts 10m (km/h)"]),    2),
                round(float(row["Wind Direction 180m (°)"]),  1),
                round(float(row["Wind Speed 180m (km/h)"]),   2),
                round(float(row["Rain (mm)"]),                3),
                round(float(row["Relative Humidity 2m (%)"]), 1),
            ])
            for cell in ws_loc[ws_loc.max_row]:
                data_style(cell, even_row=even)

        for i, w in enumerate(col_widths, 1):
            ws_loc.column_dimensions[get_column_letter(i)].width = w
        ws_loc.row_dimensions[1].height = 20
        ws_loc.freeze_panes = "A2"

    #  Summary sheet 
    ws_sum = wb.create_sheet("Summary")
    ws_sum.append(["Generated at", datetime.now().strftime("%Y-%m-%d %H:%M")])
    ws_sum.append(["Total rows",   len(df)])
    ws_sum.append(["Past days",    31])
    ws_sum.append(["Date from",    df["Date & Time"].min().strftime("%Y-%m-%d %H:%M")])
    ws_sum.append(["Date to",      df["Date & Time"].max().strftime("%Y-%m-%d %H:%M")])
    ws_sum.append(["Source",       "Open-Meteo API (open-meteo.com)"])
    ws_sum.append([])
    ws_sum.append(["Location", "Latitude", "Longitude", "Rows"])
    for loc in LOCATIONS:
        count = len(df[df["Location"] == loc["name"]])
        ws_sum.append([loc["name"], loc["lat"], loc["lon"], count])

    for row in ws_sum.iter_rows():
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            if cell.column == 1:
                cell.font = Font(name="Arial", size=10, bold=True)
    ws_sum.column_dimensions["A"].width = 14
    ws_sum.column_dimensions["B"].width = 14
    ws_sum.column_dimensions["C"].width = 14
    ws_sum.column_dimensions["D"].width = 10

    wb.save(path)
    print(f"Saved to {path}")

# Run — load existing, merge, save

existing_df = load_existing_data(EXCEL_FILE)
combined_df = merge_and_deduplicate(existing_df, new_df)
save_to_excel(combined_df, EXCEL_FILE)

# Dated backup
today       = datetime.now().strftime("%Y%m%d_%H%M")
backup_path = OUTPUT_DIR / f"weather_forecast_backup_{today}.xlsx"
save_to_excel(combined_df, backup_path)
print(f"Backup saved to {backup_path}")

